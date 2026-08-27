"""Template Analyzer Module.

Introspects PowerPoint presentation templates (.pptx) using python-pptx to extract
slide masters, slide layouts, shape placeholders, dimensions, and structural signatures.
Emits target-agnostic TemplateMetadata.
"""

from __future__ import annotations
import hashlib
import logging
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, Dict, List, Optional, Union

from pptx import Presentation as PPTXPresentation

logger = logging.getLogger(__name__)


@dataclass(slots=True)
class PlaceholderMetadata:
    """Introspected metadata for an individual slide placeholder shape.

    Attributes:
        index: Placeholder shape index within the layout.
        ph_type: String representation of PP_PLACEHOLDER type (e.g. 'TITLE', 'BODY', 'PICTURE').
        name: Shape name.
        left_inches: Shape X position in inches.
        top_inches: Shape Y position in inches.
        width_inches: Shape width in inches.
        height_inches: Shape height in inches.
    """

    index: int
    ph_type: str
    name: str
    left_inches: float
    top_inches: float
    width_inches: float
    height_inches: float


@dataclass(slots=True)
class UsableContentBounds:
    """Calculated usable content area boundaries for a slide layout.

    Attributes:
        top_inches: Content Y start position in inches.
        left_inches: Content X start position in inches.
        width_inches: Content width in inches.
        height_inches: Maximum allowed content height budget in inches.
        header_bottom_inches: Bottom coordinate of title/header region.
        footer_top_inches: Top coordinate of footer/page number region.
    """

    top_inches: float
    left_inches: float
    width_inches: float
    height_inches: float
    header_bottom_inches: float
    footer_top_inches: float


@dataclass(slots=True)
class LayoutMetadata:
    """Introspected metadata and structural signature for a single slide layout.

    Attributes:
        layout_index: Index of layout within presentation slide_layouts array.
        name: Layout name.
        placeholders: List of PlaceholderMetadata instances.
        has_title: True if layout contains a title placeholder.
        has_body: True if layout contains at least one body/content placeholder.
        has_picture: True if layout contains a picture placeholder.
        has_table: True if layout contains a table placeholder.
        body_count: Count of body content placeholders.
        brand_role: Optional brand role classification ('BRAND_COVER', 'BRAND_DISCLAIMER', etc.).
    """

    layout_index: int
    name: str
    placeholders: List[PlaceholderMetadata] = field(default_factory=list)
    has_title: bool = False
    has_body: bool = False
    has_picture: bool = False
    has_table: bool = False
    body_count: int = 0
    brand_role: Optional[str] = None

    def get_usable_content_bounds(
        self,
        template_meta: TemplateMetadata,
        header_top_margin: float = 1.0,
        footer_bottom_margin: float = 0.6,
        padding: float = 0.40,
    ) -> UsableContentBounds:
        """Calculate template-aware usable content region boundaries.

        Dynamically derives top, bottom (footer top), left, width, and max usable height bounds
        from layout placeholders and template slide dimensions without enforcing hard minimums
        that claim more space than physically exists above reserved boundaries.
        """
        s_height = template_meta.slide_height_inches
        s_width = template_meta.slide_width_inches

        # 1. Determine Header / Title Bottom Boundary
        header_bottom = header_top_margin
        title_phs = [p for p in self.placeholders if p.ph_type in ("TITLE", "CENTER_TITLE")]
        if title_phs:
            t_ph = title_phs[0]
            if t_ph.top_inches > 0 and t_ph.height_inches > 0:
                header_bottom = max(header_bottom, t_ph.top_inches + t_ph.height_inches)

        # 2. Determine Footer / Reserved Bottom Boundary
        footer_top = s_height - footer_bottom_margin
        footer_phs = [p for p in self.placeholders if p.ph_type in ("FOOTER", "SLIDE_NUMBER", "DATE")]
        if footer_phs:
            f_min_top = min(p.top_inches for p in footer_phs if p.top_inches > 0)
            if f_min_top > 0:
                footer_top = min(footer_top, f_min_top)

        # 3. Determine Body Placeholder / Content Region Bounds
        body_phs = [p for p in self.placeholders if p.ph_type in ("BODY", "OBJECT", "CONTENT", "TEXT")]
        if body_phs:
            b_ph = body_phs[0]
            top_in = max(b_ph.top_inches, header_bottom + padding)
            left_in = b_ph.left_inches if b_ph.left_inches > 0 else 0.8
            width_in = b_ph.width_inches if b_ph.width_inches > 0 else (s_width - left_in - 0.8)
            b_bottom = b_ph.top_inches + b_ph.height_inches if b_ph.height_inches > 0 else footer_top
        else:
            top_in = header_bottom + padding
            left_in = 0.8
            width_in = s_width - 1.6
            b_bottom = footer_top

        # Absolute physical upper boundary for content bottom (never cross footer or slide bottom)
        usable_bottom = min(footer_top, s_height - footer_bottom_margin)
        if b_bottom > top_in:
            usable_bottom = min(usable_bottom, b_bottom)

        # Usable height MUST NOT claim more space than exists above usable_bottom
        raw_available = usable_bottom - top_in - padding
        usable_height = max(0.2, raw_available)
        if top_in + usable_height > usable_bottom:
            usable_height = max(0.1, usable_bottom - top_in)

        return UsableContentBounds(
            top_inches=top_in,
            left_inches=left_in,
            width_inches=width_in,
            height_inches=usable_height,
            header_bottom_inches=header_bottom,
            footer_top_inches=footer_top,
        )


@dataclass(slots=True)
class TemplateMetadata:
    """Root introspected template metadata schema exported by TemplateAnalyzer.

    Attributes:
        file_path: Original template file path string.
        file_hash: SHA256 hex digest of the template file content.
        file_mtime: File modification timestamp.
        layouts: List of introspected LayoutMetadata objects (Layout Catalog).
        slide_width_inches: Presentation canvas width in inches.
        slide_height_inches: Presentation canvas height in inches.
        master_count: Total number of Slide Masters in template.
        slide_masters: List of Slide Master names.
        brand_slides: Map of detected brand roles to pre-existing slide info (Brand Slide Catalog).
        total_template_slides: Total pre-existing slides in template file.
    """

    file_path: str
    file_hash: str
    file_mtime: float
    layouts: List[LayoutMetadata] = field(default_factory=list)
    slide_width_inches: float = 10.0
    slide_height_inches: float = 7.5
    master_count: int = 1
    slide_masters: List[str] = field(default_factory=list)
    brand_slides: Dict[str, Dict[str, Any]] = field(default_factory=dict)
    total_template_slides: int = 0

    def to_dict(self) -> Dict[str, Any]:
        """Serialize metadata object to a plain dictionary for JSON caching."""
        return {
            "file_path": self.file_path,
            "file_hash": self.file_hash,
            "file_mtime": self.file_mtime,
            "slide_width_inches": self.slide_width_inches,
            "slide_height_inches": self.slide_height_inches,
            "master_count": self.master_count,
            "slide_masters": self.slide_masters,
            "brand_slides": self.brand_slides,
            "total_template_slides": self.total_template_slides,
            "layouts": [
                {
                    "layout_index": l.layout_index,
                    "name": l.name,
                    "has_title": l.has_title,
                    "has_body": l.has_body,
                    "has_picture": l.has_picture,
                    "has_table": l.has_table,
                    "body_count": l.body_count,
                    "brand_role": l.brand_role,
                    "placeholders": [
                        {
                            "index": p.index,
                            "ph_type": p.ph_type,
                            "name": p.name,
                            "left_inches": p.left_inches,
                            "top_inches": p.top_inches,
                            "width_inches": p.width_inches,
                            "height_inches": p.height_inches,
                        }
                        for p in l.placeholders
                    ],
                }
                for l in self.layouts
            ],
        }

    @classmethod
    def from_dict(cls, data: Dict[str, Any]) -> TemplateMetadata:
        """Instantiate TemplateMetadata from a cached dictionary."""
        layouts = []
        for l_data in data.get("layouts", []):
            phs = [PlaceholderMetadata(**p) for p in l_data.get("placeholders", [])]
            layouts.append(
                LayoutMetadata(
                    layout_index=l_data["layout_index"],
                    name=l_data["name"],
                    placeholders=phs,
                    has_title=l_data.get("has_title", False),
                    has_body=l_data.get("has_body", False),
                    has_picture=l_data.get("has_picture", False),
                    has_table=l_data.get("has_table", False),
                    body_count=l_data.get("body_count", 0),
                    brand_role=l_data.get("brand_role"),
                )
            )
        return cls(
            file_path=data["file_path"],
            file_hash=data["file_hash"],
            file_mtime=data["file_mtime"],
            slide_width_inches=data.get("slide_width_inches", 10.0),
            slide_height_inches=data.get("slide_height_inches", 7.5),
            master_count=data.get("master_count", 1),
            slide_masters=data.get("slide_masters", []),
            brand_slides=data.get("brand_slides", {}),
            total_template_slides=data.get("total_template_slides", 0),
            layouts=layouts,
        )


class TemplateAnalyzer:
    """Analyzer introspecting PPTX templates and outputting TemplateMetadata."""

    def analyze(self, pptx_path: Union[str, Path]) -> TemplateMetadata:
        """Introspect a PowerPoint presentation template (.pptx).

        Args:
            pptx_path: Path to target PPTX file.

        Returns:
            TemplateMetadata instance.

        Raises:
            FileNotFoundError: If pptx_path does not exist.
        """
        path = Path(pptx_path)
        if not path.is_file():
            raise FileNotFoundError(f"PowerPoint template not found: {pptx_path}")

        file_bytes = path.read_bytes()
        file_hash = hashlib.sha256(file_bytes).hexdigest()
        file_mtime = path.stat().st_mtime

        prs = PPTXPresentation(str(path))
        width_in = float(prs.slide_width.inches) if hasattr(prs, "slide_width") else 10.0
        height_in = float(prs.slide_height.inches) if hasattr(prs, "slide_height") else 7.5

        slide_masters_names = [getattr(master, "name", f"Master {idx}") for idx, master in enumerate(prs.slide_masters)]

        layouts_meta: List[LayoutMetadata] = []
        brand_slides_map: Dict[str, Dict[str, Any]] = {}

        # 1. Build Layout Catalog (for dynamic AI slides)
        for idx, layout in enumerate(prs.slide_layouts):
            phs_meta: List[PlaceholderMetadata] = []
            has_title = False
            has_body = False
            has_picture = False
            has_table = False
            body_count = 0

            for ph in layout.placeholders:
                ph_type_str = str(ph.placeholder_format.type).split(".")[-1].split(" ")[0].upper()
                
                left_in = float(ph.left.inches) if hasattr(ph, "left") and ph.left else 0.0
                top_in = float(ph.top.inches) if hasattr(ph, "top") and ph.top else 0.0
                w_in = float(ph.width.inches) if hasattr(ph, "width") and ph.width else 0.0
                h_in = float(ph.height.inches) if hasattr(ph, "height") and ph.height else 0.0

                if ph_type_str in ("TITLE", "CENTER_TITLE"):
                    has_title = True
                elif ph_type_str in ("BODY", "OBJECT", "CONTENT", "TEXT"):
                    has_body = True
                    body_count += 1
                elif ph_type_str == "PICTURE":
                    has_picture = True
                elif ph_type_str == "TABLE":
                    has_table = True

                phs_meta.append(
                    PlaceholderMetadata(
                        index=int(ph.placeholder_format.idx),
                        ph_type=ph_type_str,
                        name=ph.name,
                        left_inches=left_in,
                        top_inches=top_in,
                        width_inches=w_in,
                        height_inches=h_in,
                    )
                )

            layouts_meta.append(
                LayoutMetadata(
                    layout_index=idx,
                    name=layout.name,
                    placeholders=phs_meta,
                    has_title=has_title,
                    has_body=has_body,
                    has_picture=has_picture,
                    has_table=has_table,
                    body_count=body_count,
                )
            )

        # 2. Build Brand Slide Catalog (inspect pre-existing template slides using multi-heuristics)
        total_slides_in_template = len(prs.slides)

        for s_idx, s in enumerate(prs.slides):
            s_title = ""
            full_text_corpus = []
            img_count = 0
            has_logo = False
            has_footer = False
            has_smartart = False
            editable_phs = []

            for shp in s.shapes:
                if shp.shape_type == 13:  # Picture shape
                    img_count += 1
                if "logo" in shp.name.lower() or "brand" in shp.name.lower():
                    has_logo = True
                if "smartart" in shp.name.lower() or "diagram" in shp.name.lower() or getattr(shp, "has_smart_art", False):
                    has_smartart = True

                if shp.is_placeholder:
                    ph_t = str(shp.placeholder_format.type).split(".")[-1].split(" ")[0].upper()
                    if ph_t in ("TITLE", "CENTER_TITLE", "SUBTITLE", "BODY", "OBJECT", "TEXT"):
                        editable_phs.append(ph_t)

                if shp.has_text_frame and shp.text_frame.text.strip():
                    txt = shp.text_frame.text.strip()
                    full_text_corpus.append(txt)
                    if not s_title or shp.name.startswith("Title"):
                        s_title = txt

                    if "copyright" in txt.lower() or "iso 27001" in txt.lower():
                        has_footer = True

            combined_text = " ".join(full_text_corpus).lower()
            detected_role: Optional[str] = None

            # Multi-heuristic brand role classification
            if "disclaimer" in combined_text or "proprietary to molecular connections" in combined_text:
                detected_role = "BRAND_DISCLAIMER"
            elif "about company" in combined_text or ("machine learning company" in combined_text and "multilingual" in combined_text):
                detected_role = "BRAND_ABOUT_COMPANY"
            elif "case study" in combined_text or "lorem ipsum" in combined_text:
                if "BRAND_CASE_STUDY" not in brand_slides_map:
                    detected_role = "BRAND_CASE_STUDY"
            elif "thank you" in combined_text or "thanks" in combined_text or (s_idx == total_slides_in_template - 1 and img_count > 0):
                detected_role = "BRAND_THANK_YOU"
            elif s_idx == 0 or "title here" in combined_text or "welcome" in combined_text:
                detected_role = "BRAND_COVER"

            if detected_role and detected_role not in brand_slides_map:
                l_id = 0
                if hasattr(s.slide_layout, "slide_layout_id"):
                    l_id = int(s.slide_layout.slide_layout_id)
                l_name = getattr(s.slide_layout, "name", "Brand Layout")

                brand_slides_map[detected_role] = {
                    "role": detected_role,
                    "slide_index": s_idx,
                    "layout_index": l_id,
                    "layout_name": l_name,
                    "title": s_title or detected_role.replace("BRAND_", "").replace("_", " ").title(),
                    "shapes_count": len(s.shapes),
                    "images_count": img_count,
                    "logo_detected": has_logo,
                    "footer_detected": has_footer,
                    "smartart_detected": has_smartart,
                    "relationships_preserved": True,
                    "theme_preserved": True,
                    "editable_placeholders": editable_phs,
                }

        return TemplateMetadata(
            file_path=str(path.resolve()),
            file_hash=file_hash,
            file_mtime=file_mtime,
            layouts=layouts_meta,
            slide_width_inches=width_in,
            slide_height_inches=height_in,
            master_count=len(prs.slide_masters),
            slide_masters=slide_masters_names,
            brand_slides=brand_slides_map,
            total_template_slides=total_slides_in_template,
        )
